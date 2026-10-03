#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成选手版 SDK API 参考：**每个选手程序一份**。

    python3 scripts/gen_sdk_api_doc.py            # 每个选手程序一份 .md + .docx
    python3 scripts/gen_sdk_api_doc.py --check    # 只检查分类有没有漏，不写文件

每份只收录**那一个程序真正用过的 API**，不是全部 80 个，也不是四个程序的并集。
用过哪些是从程序源码里用 AST 扫出来的，程序改了重新跑一遍就同步——不手工
维护名单，手工名单迟早跟代码对不上。

**条目按程序里出现的先后排，并标出首次用到的行号**：对着代码从上往下读，
文档顺序正好能对上，不用在分类目录里来回找。

签名和说明同样从 capabilities.py 读。分类表是人工定的（机器分不出该放哪类），
扫出来的 API 如果没进分类表会直接报错，而不是悄悄漏掉。
"""
import argparse
import ast
import glob
import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'src', 'contest_sdk', 'contest_sdk', 'capabilities.py')
PROGS = os.path.join(ROOT, 'contestant_sim', '*_lite.py')
OUTDIR = os.path.join(ROOT, 'contestant_sim', 'API参考')

CATEGORIES = [
    ('程序入口', ['run', 'spacing_m']),
    ('起飞 / 降落 / 返航', ['takeoff', 'return_home', 'own_pad']),
    ('航线飞行', ['fly_route', 'goto_world', 'hold_at', 'set_agl', 'step_forward']),
    ('编队', ['lead_formation', 'follow_formation']),
    ('跨机协同', ['open_inbox', 'wait_event', 'wait_any_event',
                  'send_to_teammate', 'yield_spot']),
    ('视觉：识别与对准', ['search_along', 'patrol', 'aim_at']),
    ('拍照', ['snapshot', 'PHOTO_DIR']),
    ('抓放与发射', ['grip', 'fetch_from', 'release_at', 'shoot']),
    ('声光播报', ['announce']),
    ('日志', ['progress']),
]

PITFALLS = {
    'open_inbox': '可靠事件通道是“先回 ACK 再查处理函数”，没注册的事件会被确认后丢弃。'
                  '任务一开始就把所有事件注册全，别等用到了再注册。',
    'aim_at': 'camera=\'down\' 会自动改走 precision_servo 闭环。前视那套几何把画面纵轴'
              '当成世界的高低，下视时纵轴其实是机体前后方向，照搬永远收敛不了。',
    'goto_world': 'direct=True 会关掉避障，只在算过余量的航段上开（tools/check_route.py）。'
                  '航线上的避障本身是考核点，不能为了快绕过去。',
    'snapshot': '不要给拍照配声光事件：声光事件是固定枚举，自造事件名会直接抛 ValueError '
                '把整个任务打断。',
    'return_home': '自带落点核对：落在起降点上才播 sound、才算数。任务流程里会降落好几次'
                   '（取器材、放器材、回家），“降落动作完成”本身说明不了任务结束。',
    'lead_formation': 'disband_at 要显式给。不给的话默认判据是“长机飞回自己起飞点上空”，'
                      '而起飞点不一定在航线上，编队可能在半路散掉。',
    'shoot': '连发时 sound 只播一次，播报点在第一发之前。',
    'PHOTO_DIR': 'snapshot() 的存图目录，用实例属性覆盖即可：sdk.PHOTO_DIR = \'/logs/xxx\'。',
    'spacing_m': 'run() 把命令行 --spacing 存成这个属性，任务函数直接读。',
}


def used_apis():
    """{程序文件名: [(API 名, 首次用到的行号), ...]}，按行号排序。"""
    out = {}
    for f in sorted(glob.glob(PROGS)):
        first = {}
        for n in ast.walk(ast.parse(io.open(f, encoding='utf-8').read())):
            if isinstance(n, ast.Attribute) and isinstance(n.value, ast.Name) \
                    and n.value.id in ('sdk', 'DroneSDK'):
                if n.attr not in first or n.lineno < first[n.attr]:
                    first[n.attr] = n.lineno
        out[os.path.basename(f)] = sorted(first.items(), key=lambda kv: kv[1])
    return out


def param_gaps(api):
    """每个在用 API 的**每个参数**都必须在 docstring 的 Args 里有说明。

    返回 [(方法名, [缺说明的参数])]。这是一道闸：参数说明不全就不出文档。
    2026-10-01 第一次量的时候 96 个参数里缺 64 个（67%），全靠这个脚本逐个
    点名才补齐——没有闸的话下次加参数又会忘。
    """
    cls = next(n for n in ast.parse(io.open(SRC, encoding='utf-8').read()).body
               if isinstance(n, ast.ClassDef) and n.name == 'DroneSDK')
    fn = {n.name: n for n in cls.body
          if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))}
    gaps = []
    for name, info in api.items():
        if name not in fn:
            continue                      # 属性，没有参数
        a = fn[name].args
        ps = [x.arg for x in a.posonlyargs + a.args + a.kwonlyargs if x.arg != 'self']
        if a.vararg:
            ps.append(a.vararg.arg)
        if a.kwarg:
            ps.append(a.kwarg.arg)
        doc = {d.split(':')[0].strip().lstrip('*') for d in info['args']}
        lack = [x for x in ps if x not in doc]
        if lack:
            gaps.append((name, lack))
    return gaps


def category_of(name):
    for cat, names in CATEGORIES:
        if name in names:
            return cat
    return '其它'



def collect():
    cls = next(n for n in ast.parse(io.open(SRC, encoding='utf-8').read()).body
               if isinstance(n, ast.ClassDef) and n.name == 'DroneSDK')
    api = {}
    for n in cls.body:
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and not n.name.startswith('_'):
            doc = ast.get_docstring(n) or ''
            para, args = [], []
            lines = doc.splitlines()
            i = 0
            while i < len(lines) and (lines[i].strip() or not para):
                if lines[i].strip():
                    para.append(lines[i].strip())
                i += 1
            summary = ''.join(para)
            if 'Args:' in doc:
                # 参数说明经常跨好几行。**按缩进判断**：缩进等于基准的那一行
                # 是新参数，更深的是上一条的续行，要接上去——否则一条说明会被
                # 拆成好几个条目，读起来是断的。
                blk = doc.split('Args:', 1)[1].splitlines()[1:]
                base = None
                for line in blk:
                    if not line.strip():
                        if args:
                            break
                        continue
                    ind = len(line) - len(line.lstrip())
                    if base is None:
                        base = ind
                    if ind <= base and line.strip().endswith(':') and ' ' not in line.strip():
                        break                      # 碰到 Returns:/Raises: 这类下一节
                    if ind > base and args:
                        args[-1] = args[-1].rstrip() + ' ' + line.strip()
                    else:
                        args.append(line.strip())
            sig = ast.unparse(n.args).replace('self, ', '').replace('self', '')
            api[n.name] = {'sig': f'{n.name}({sig})', 'summary': summary, 'args': args,
                           'static': any(isinstance(d, ast.Name) and d.id == 'staticmethod'
                                         for d in n.decorator_list)}
    for n in ast.walk(cls):      # 类属性。两种写法都要认：
        #   PHOTO_DIR = '...'        -> ast.Assign
        #   PHOTO_DIR: str = '...'   -> ast.AnnAssign（加类型标注之后是这种）
        tgt = None
        if isinstance(n, ast.Assign) and isinstance(n.targets[0], ast.Name):
            tgt = n.targets[0].id
        elif isinstance(n, ast.AnnAssign) and isinstance(n.target, ast.Name):
            tgt = n.target.id
        if tgt:
            api.setdefault(tgt, {'sig': tgt, 'summary': '',
                                 'args': [], 'static': False})
    for n in ast.walk(cls):      # 实例属性（spacing_m）
        if isinstance(n, ast.Attribute) and isinstance(n.ctx, ast.Store) \
                and isinstance(n.value, ast.Name) and n.value.id in ('self', 'sdk'):
            api.setdefault(n.attr, {'sig': n.attr, 'summary': '', 'args': [], 'static': False})
    return api


def _rich(par, text, base_size=None):
    """把 docstring 里的 Markdown 行内标记渲染成 Word 的格式。

    源码 docstring 用 **粗体** 和 `等宽` 标记重点，直接塞进 docx 就是一堆
    星号和反引号，在 Word 里很难看。这里按标记切段，分别设成粗体 / Consolas。
    """
    import re
    from docx.shared import Pt
    for seg in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**'):
            r = par.add_run(seg[2:-2]); r.bold = True
        elif seg.startswith('`') and seg.endswith('`'):
            r = par.add_run(seg[1:-1]); r.font.name = 'Consolas'
        else:
            r = par.add_run(seg)
        if base_size:
            r.font.size = Pt(base_size)


def _plain(text):
    """去掉 Markdown 标记，给表格单元格用（单元格里不做富文本）。"""
    import re
    return re.sub(r'\*\*([^*]+)\*\*', r'\1', text).replace('`', '')


def build_md(api, prog, items):
    # 属性（PHOTO_DIR / spacing_m）没有 docstring，它们的说明写在 PITFALLS 里。
    # 这时该当正文用，不该标成红色警告——那是给"有说明、另外还有坑"的方法留的。
    """一个程序一份 Markdown，条目按程序里出现的先后排。"""
    L = [f'# {prog} 用到的 SDK API\n',
         f'> 这个程序用到 **{len(items)} 个** API（SDK 全部能力有 80 个，'
         f'其余是底层/备用接口，这个程序没用到）。\n'
         f'> 条目按**程序里出现的先后**排，括号里是首次用到的行号——'
         f'对着 `{prog}` 从上往下读，顺序能对上。\n'
         f'> 本文档由 `scripts/gen_sdk_api_doc.py` 自动生成，**不要手改**。\n',
         '\n## 速查\n', '| 行 | API | 分类 | 做什么 |', '|---|---|---|---|']
    for name, line in items:
        desc = _plain(api[name]['summary'] or PITFALLS.get(name, ''))[:52]
        L.append(f'| {line} | [`{name}`](#{name.lower()}) | {category_of(name)} | {desc} |')
    L.append('\n## 逐个说明\n')
    for name, line in items:
        info = api[name]
        L.append(f'<a id="{name.lower()}"></a>')
        L.append(f'### `{info["sig"]}`\n')
        L.append(f'*{prog} 第 {line} 行首次用到 · {category_of(name)}*'
                 + ('  ·  *staticmethod*' if info['static'] else '') + '\n')
        body = info['summary'] or (PITFALLS.get(name, '') if not info['args'] else '')
        if body:
            L.append(body + '\n')
        if info['args']:
            L.append('**参数**\n')
            for a in info['args']:
                L.append(f'- {a}')
            L.append('')
        if name in PITFALLS and body != PITFALLS[name]:
            L.append(f'> ⚠️ {PITFALLS[name]}\n')
    return '\n'.join(L) + '\n'


def build_docx(api, prog, items):
    """同样内容的 .docx。docstring 的 **粗体** / `等宽` 转成 Word 真实格式。"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    d = Document()
    d.styles['Normal'].font.size = Pt(10.5)
    d.styles['Normal'].font.name = '微软雅黑'

    d.add_heading(f'{prog} 用到的 SDK API', 0)
    p = d.add_paragraph()
    p.add_run(f'这个程序用到 {len(items)} 个 API。').bold = True
    p.add_run('SDK 全部能力有 80 个，其余是底层/备用接口，这个程序没用到。'
              '条目按程序里出现的先后排，括号里是首次用到的行号——对着源码'
              '从上往下读，顺序能对上。本文档自动生成，不要手改。')

    d.add_heading('速查', level=1)
    t = d.add_table(rows=1, cols=4)
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(('行', 'API', '分类', '做什么')):
        t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for name, line in items:
        r = t.add_row().cells
        r[0].text = str(line)
        r[1].paragraphs[0].add_run(name).font.name = 'Consolas'
        r[2].text = category_of(name)
        r[3].text = _plain(api[name]['summary'] or PITFALLS.get(name, ''))[:52]

    d.add_page_break()
    d.add_heading('逐个说明', level=1)
    for name, line in items:
        info = api[name]
        h = d.add_paragraph()
        r = h.add_run(info['sig'])
        r.bold = True
        r.font.name = 'Consolas'
        r.font.size = Pt(11)
        meta = d.add_paragraph()
        mr = meta.add_run(f'第 {line} 行首次用到 · {category_of(name)}'
                          + ('  ·  staticmethod' if info['static'] else ''))
        mr.italic = True
        mr.font.size = Pt(9)
        body = info['summary'] or (PITFALLS.get(name, '') if not info['args'] else '')
        if body:
            _rich(d.add_paragraph(), body)
        if info['args']:
            d.add_paragraph().add_run('参数').bold = True
            for a in info['args']:
                ap = d.add_paragraph(style='List Bullet')
                ap.paragraph_format.left_indent = Pt(24)
                _rich(ap, a)
        if name in PITFALLS and body != PITFALLS[name]:
            w = d.add_paragraph()
            w.add_run('⚠ ')
            _rich(w, PITFALLS[name])
            for rr in w.runs:
                rr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    return d


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--check', action='store_true')
    args = ap.parse_args()

    api, used = collect(), used_apis()
    allnames = {n for items in used.values() for n, _ in items}
    listed = {m for _, ms in CATEGORIES for m in ms}
    missing = sorted(allnames - listed)
    unknown = sorted(allnames - set(api))
    stale = sorted(listed - allnames)
    if missing or unknown:
        if missing:
            print(f'★NG  选手程序用了但没进分类表：{missing}', file=sys.stderr)
        if unknown:
            print(f'★NG  选手程序用了 SDK 里不存在的名字：{unknown}', file=sys.stderr)
        return 1
    if stale:
        print(f'注意  分类表里这些已经没有程序在用，已跳过：{stale}', file=sys.stderr)
    gaps = [(n, l) for n, l in param_gaps(api) if n in allnames]
    if gaps:
        for n, l in gaps:
            print(f'★NG  {n}() 这些参数在 docstring 的 Args 里没有说明：'
                  f'{"、".join(l)}', file=sys.stderr)
        print('     去 capabilities.py 对应方法的 docstring 里补上 Args。',
              file=sys.stderr)
        return 1
    if args.check:
        n_par = sum(len(api[x]['args']) for x in allnames)
        print(f'OK   {len(allnames)} 个在用 API 全部有分类，{n_par} 个参数全部有说明')
        return 0
    os.makedirs(OUTDIR, exist_ok=True)
    for prog, items in used.items():
        stem = prog[:-3]
        io.open(os.path.join(OUTDIR, stem + '.md'), 'w',
                encoding='utf-8').write(build_md(api, prog, items))
        build_docx(api, prog, items).save(os.path.join(OUTDIR, stem + '.docx'))
        print(f'已生成 API参考/{stem}.md 和 .docx（{len(items)} 个 API）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
